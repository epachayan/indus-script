import csv,glob,os,random,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
R=sys.argv[1]; SEED=20261010; rng=random.Random(SEED)
def rows(pattern,hdr):
    out=[]
    for f in sorted(glob.glob(pattern)):
        for r in csv.reader(open(f,encoding='utf-8')):
            if not r or r[0]==hdr[0]: continue
            out.append((os.path.basename(f),dict(zip(hdr,r))))
    return out
mk=rows(R+'/transcriptions/mackay1938_seal_table/*.csv',
 'no,type,dims_in,material,block,house,room,street,level_ft,field_no,pdf_page'.split(','))
ms=rows(R+'/transcriptions/marshall1931_seal_table/*.csv',
 'plate_no,size_in,level_ft,level_in,type,material,area,serial'.split(','))
pc=list(csv.DictReader(open(R+'/transcriptions/mackay1938_plates/photo_codings.csv',encoding='utf-8')))
print(len(mk),len(ms),len(pc))
wb=Workbook(); hf=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='305496')
inp=PatternFill('solid',fgColor='FFF2CC')
def sheet(ws,where_cols,fields,data):
    hdr=['#']+where_cols+fields+['Row correct? (Y/N)','Wrong field(s)','Correct value as printed']
    ws.append(hdr)
    for c in ws[1]: c.font=hf; c.fill=fill; c.alignment=Alignment(wrap_text=True)
    dv=DataValidation(type='list',formula1='"Y,N"',allow_blank=True); ws.add_data_validation(dv)
    for i,(w,r) in enumerate(data,1):
        ws.append([i]+w+[r.get(f,'') for f in fields]+['','',''])
        n=ws.max_row; k=len(hdr)
        dv.add(f'{ws.cell(n,k-2).coordinate}')
        for j in range(k-2,k+1): ws.cell(n,j).fill=inp
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=max(8,min(28,max(len(str(c.value or '')) for c in col)+2))
    ws.freeze_panes='B2'
    return len(data)
s1=sorted(rng.sample(mk,50),key=lambda x:int(x[1]['pdf_page']))
ws=wb.active; ws.title='Mackay 1938 (50)'
f1='no,type,dims_in,material,block,house,room,street,level_ft,field_no'.split(',')
sheet(ws,['Printed page','Scan page (chapter XI extract)'],f1,
 [([int(r['pdf_page'])+324,int(r['pdf_page'])],r) for _,r in s1])
s2=sorted(rng.sample(ms,50),key=lambda x:x[0])
ws=wb.create_sheet('Marshall 1931 (50)')
def loc(fn): p=fn[1:4]; c='left' if fn[4]=='L' else 'right'; return [int(p),int(p)-346,c]
sheet(ws,['Printed page','PDF page','Column'],'plate_no,size_in,level_ft,level_in,type,material,area,serial'.split(','),
 [(loc(fn),r) for fn,r in s2])
s3=rng.sample(pc,30)
order=['LXXXII','LXXXIII','LXXXIV','LXXXV','LXXXVI','LXXXVII','LXXXVIII','LXXXIX','XCIV','XCV','XCVI','XCVII','XCVIII','XCIX']
s3=sorted(s3,key=lambda r:(order.index(r['plate']) if r['plate'] in order else 99,int(''.join(ch for ch in r['mackay_no'] if ch.isdigit()) or 0)))
ws=wb.create_sheet('Plate codings (30)')
sheet(ws,['Plate','Seal no.'],'n_signs,ends_jar,has_fish,has_arrow,numeral_group,confidence,note'.split(','),
 [([r['plate'],r['mackay_no']],r) for r in s3])
ws=wb.create_sheet('Summary',0)
ws.append(['Spot-check of AI transcriptions against the scans']); ws['A1'].font=Font(bold=True,size=14)
ws.append([f'Random sample, seed {SEED} (Python random.Random). Re-running make_sheet.py reproduces the same rows.'])
ws.append(['Mark each row Y if every field matches the printed page, N otherwise; note the wrong field and the printed value.'])
ws.append([])
ws.append(['Set','Population','Sampled','Checked','Errors (rows with N)','Row error rate'])
for c in ws[5]: c.font=hf; c.fill=fill
for i,(nm,pop,n) in enumerate([('Mackay 1938 (50)',len(mk),50),('Marshall 1931 (50)',len(ms),50),('Plate codings (30)',len(pc),30)],6):
    k=chr(ord('A')+len(f1)+2+1+0)
    ws.append([nm,pop,n,f'=COUNTIF(\'{nm}\'!{{c}}:{{c}},"Y")+COUNTIF(\'{nm}\'!{{c}}:{{c}},"N")',f'=COUNTIF(\'{nm}\'!{{c}}:{{c}},"N")',f'=IF(D{i}=0,"",E{i}/D{i})'])
# fix column letters per sheet
for i,nm in enumerate(['Mackay 1938 (50)','Marshall 1931 (50)','Plate codings (30)'],6):
    s=wb[nm]; col=[c.column_letter for c in s[1] if c.value=='Row correct? (Y/N)'][0]
    for x in 'DE': ws[f'{x}{i}']=ws[f'{x}{i}'].value.replace('{c}',col)
    ws[f'F{i}'].number_format='0.0%'
for c,w in zip('ABCDEF',[24,12,10,10,20,16]): ws.column_dimensions[c].width=w
ws.append([]); ws.append(['Notes']); ws['A11'].font=Font(bold=True)
for t in ['Mackay: printed page = scan page + 324 (TVA_BOK_0009130, ch. XI "Tabulation of Seals"). Ditto marks were resolved to the value above.',
          'Marshall: archive.org in.ernet.dli.2015.107493, pp. 402-405; PDF page = printed page - 346. level_ft/level_in = depth below surface.',
          'Plates: Mackay vol. II, full-resolution IIIF scans (in.gov.ignca.9954). Check sign count and jar ending first; those are what the analysis uses.',
          'A row counts as an error if any field is wrong. Also note anything the transcription missed entirely.']:
    ws.append([t])
wb.save(os.path.expanduser('~/spotcheck/spotcheck_sample.xlsx'))
